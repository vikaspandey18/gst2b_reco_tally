import io
import unittest
from pathlib import Path
import openpyxl

from app.controllers.parser import extract_pan, clean_str, clean_invoice_no, clean_num, parse_excel_records
from app.controllers.reco_engine import run_reconciliation, normalize_inv_str, filter_duplicates
from app.controllers.excel_generator import generate_reco_excel

class TestGstReconciliation(unittest.TestCase):

    def test_pan_extraction(self):
        """Test that 15-digit GSTIN and 10-digit PAN both resolve to the correct 10-digit PAN."""
        # 15-char GSTIN
        self.assertEqual(extract_pan("27AAICT8273F1ZY"), "AAICT8273F")
        self.assertEqual(extract_pan("08AAECP9085D2Z5"), "AAECP9085D")
        self.assertEqual(extract_pan("27AASCA5981H1Z2"), "AASCA5981H")
        
        # 10-char PAN
        self.assertEqual(extract_pan("AAICT8273F"), "AAICT8273F")
        self.assertEqual(extract_pan("aaecp9085d"), "AAECP9085D")
        self.assertEqual(extract_pan(" AASCA5981H "), "AASCA5981H")

    def test_matching_when_book_has_pan_and_2b_has_gstin(self):
        """Test reconciliation matches when one file has 10-digit PAN and the other has 15-digit GSTIN."""
        book_records = [{
            "period": "April 2025",
            "trade_name": "Vendor A",
            "gstin_raw": "AAICT8273F",        # PAN passed in Book
            "pan": "AAICT8273F",
            "gstin": "AAICT8273F",
            "inv_type": "Purchase",
            "reverse_charge": "No",
            "invoice_no": "INV-001",
            "reco_key": "AAICT8273F_INV-001",
            "invoice_date": "2025-04-10",
            "invoice_value": 11800.0,
            "taxable_value": 10000.0,
            "tax_rate": 18,
            "igst": 0.0,
            "cgst": 900.0,
            "sgst": 900.0,
            "cess": 0.0,
            "total_tax": 1800.0
        }]

        two_b_records = [{
            "period": "April 2025",
            "trade_name": "Vendor A Private Limited",
            "gstin_raw": "27AAICT8273F1ZY",    # Full GSTIN passed in GSTR-2B
            "pan": "AAICT8273F",
            "gstin": "AAICT8273F",
            "inv_type": "Regular",
            "reverse_charge": "No",
            "invoice_no": "INV-001",
            "reco_key": "AAICT8273F_INV-001",
            "invoice_date": "2025-04-10",
            "invoice_value": 11800.0,
            "taxable_value": 10000.0,
            "tax_rate": 18,
            "igst": 0.0,
            "cgst": 900.0,
            "sgst": 900.0,
            "cess": 0.0,
            "total_tax": 1800.0
        }]

        result = run_reconciliation(book_records, two_b_records)
        stats = result["stats"]

        self.assertEqual(stats["total_invoices"], 1)
        self.assertEqual(stats["matched_count"], 1)
        self.assertEqual(stats["mismatched_count"], 0)
        self.assertEqual(stats["total_tax_diff"], 0.0)

        # Check item remarks
        item = result["invoice_reco_items"][0]
        self.assertEqual(item["final_remarks"], "Book And 2B Match")
        self.assertEqual(item["status_x"], "Book And 2B Match")

    def test_duplicate_invoices_handling(self):
        """Test duplicate invoices are separated into duplicate_invoices list and unique invoices are used for reco."""
        raw_book_records = [
            {
                "period": "April 2025",
                "trade_name": "Vendor B",
                "gstin_raw": "27AASCA5981H1Z2",
                "pan": "AASCA5981H",
                "gstin": "AASCA5981H",
                "invoice_no": "DUP-100",
                "reco_key": "AASCA5981H_DUP-100",
                "taxable_value": 5000.0,
                "igst": 900.0,
                "cgst": 0.0,
                "sgst": 0.0,
                "cess": 0.0,
                "total_tax": 900.0
            },
            {
                # Duplicate entry in Book!
                "period": "April 2025",
                "trade_name": "Vendor B",
                "gstin_raw": "27AASCA5981H1Z2",
                "pan": "AASCA5981H",
                "gstin": "AASCA5981H",
                "invoice_no": "DUP-100",
                "reco_key": "AASCA5981H_DUP-100",
                "taxable_value": 5000.0,
                "igst": 900.0,
                "cgst": 0.0,
                "sgst": 0.0,
                "cess": 0.0,
                "total_tax": 900.0
            }
        ]

        raw_two_b_records = [
            {
                "period": "April 2025",
                "trade_name": "Vendor B",
                "gstin_raw": "27AASCA5981H1Z2",
                "pan": "AASCA5981H",
                "gstin": "AASCA5981H",
                "invoice_no": "DUP-100",
                "reco_key": "AASCA5981H_DUP-100",
                "taxable_value": 5000.0,
                "igst": 900.0,
                "cgst": 0.0,
                "sgst": 0.0,
                "cess": 0.0,
                "total_tax": 900.0
            }
        ]

        result = run_reconciliation(raw_book_records, raw_two_b_records)

        # Verify duplicate count
        self.assertEqual(len(result["duplicate_invoices"]), 1)
        dup = result["duplicate_invoices"][0]
        self.assertEqual(dup["source"], "Book")
        self.assertEqual(dup["invoice_no"], "DUP-100")
        self.assertEqual(dup["duplicate_count"], 2)

        # Unique records should have only 1 item in reconciliation
        self.assertEqual(result["stats"]["total_invoices"], 1)
        self.assertEqual(result["stats"]["matched_count"], 1)
        self.assertEqual(result["stats"]["duplicate_count"], 1)

    def test_excel_generation_with_5_sheets(self):
        """Test Excel generator creates all 5 sheets including 'Duplicate Invoices'."""
        book_records = [{
            "period": "May 2025",
            "trade_name": "Vendor C",
            "gstin_raw": "AAICT8273F",
            "pan": "AAICT8273F",
            "gstin": "AAICT8273F",
            "invoice_no": "INV-555",
            "reco_key": "AAICT8273F_INV-555",
            "taxable_value": 2000.0,
            "total_tax": 360.0,
            "igst": 360.0,
            "cgst": 0.0,
            "sgst": 0.0,
            "cess": 0.0
        }, {
            # Duplicate
            "period": "May 2025",
            "trade_name": "Vendor C",
            "gstin_raw": "AAICT8273F",
            "pan": "AAICT8273F",
            "gstin": "AAICT8273F",
            "invoice_no": "INV-555",
            "reco_key": "AAICT8273F_INV-555",
            "taxable_value": 2000.0,
            "total_tax": 360.0,
            "igst": 360.0,
            "cgst": 0.0,
            "sgst": 0.0,
            "cess": 0.0
        }]
        two_b_records = [{
            "period": "May 2025",
            "trade_name": "Vendor C",
            "gstin_raw": "27AAICT8273F1ZY",
            "pan": "AAICT8273F",
            "gstin": "AAICT8273F",
            "invoice_no": "INV-555",
            "reco_key": "AAICT8273F_INV-555",
            "taxable_value": 2000.0,
            "total_tax": 360.0,
            "igst": 360.0,
            "cgst": 0.0,
            "sgst": 0.0,
            "cess": 0.0
        }]

        reco_res = run_reconciliation(book_records, two_b_records)
        out_path = Path("storage/exports/test_5_sheets.xlsx")
        generate_reco_excel(reco_res, out_path)

        # Verify generated workbook sheets
        wb = openpyxl.load_workbook(str(out_path), data_only=True)
        expected_sheets = ["GST + Invocie No", "GSTR-2B", "Book", "GST Wise Reco", "Duplicate Invoices"]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Verify Duplicate Invoices sheet contains the duplicate
        ws_dup = wb["Duplicate Invoices"]
        self.assertEqual(ws_dup.cell(1, 1).value, "Duplicate Invoices Report")
        self.assertEqual(ws_dup.cell(4, 1).value, "Book")
        self.assertEqual(ws_dup.cell(4, 5).value, "INV-555")

if __name__ == "__main__":
    unittest.main()
