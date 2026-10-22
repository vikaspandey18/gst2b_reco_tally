import re
import io
import openpyxl
from typing import List, Dict, Any, Tuple

# Essential header aliases for GSTR-2B and Tally Book
HEADER_ALIASES = {
    "period": ["period", "return period", "month"],
    "trade_name": ["trade/legal name", "trade name", "party name", "supplier name", "particulars", "name of supplier"],
    "gstin": ["gstin of supplier", "gstin", "supplier gstin", "gst no", "party gstin", "pan", "pan no"],
    "inv_type": ["invoice type", "vch type", "voucher type", "document type", "type"],
    "reverse_charge": ["reverse charge", "rcm"],
    "invoice_no": ["invoice no", "invoice number", "vch no.", "vch no", "voucher no", "doc no", "bill no"],
    "invoice_date": ["invoice date", "date", "vch date"],
    "invoice_value": ["invoice value", "total value", "gross amount", "voucher total"],
    "taxable_value": ["taxable value", "taxable amount", "assessable value", "taxable"],
    "tax_rate": ["tax rate", "rate", "gst rate"],
    "igst": ["integrated tax", "igst", "igst amount"],
    "cgst": ["central tax", "cgst", "cgst amount"],
    "sgst": ["state/ut tax", "sgst", "sgst amount", "utgst"],
    "cess": ["cess", "cess amount"]
}

def clean_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).replace("_x000D_\n", "").replace("_x000D_", "").replace("\r", "").replace("\n", " ").strip()

def clean_invoice_no(val: Any) -> str:
    s = clean_str(val)
    return s.strip("'").strip('"').strip()

def clean_num(val: Any) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return round(float(val), 2)
    s = str(val).replace(",", "").replace("₹", "").replace("Rs", "").strip()
    try:
        return round(float(s), 2)
    except (ValueError, TypeError):
        return 0.0

def extract_pan(val: Any) -> str:
    """
    Extracts 10-character PAN number from either a 15-character GSTIN or a 10-character PAN.
    If full GSTIN is provided (e.g. 27AAICT8273F1ZY), extracts positions 3-12 (AAICT8273F).
    If 10-char PAN is provided directly, returns it in uppercase.
    """
    if not val:
        return ""
    s = clean_str(val).upper().replace(" ", "").replace("-", "")
    # Standard 15-character GSTIN
    if len(s) == 15 and re.match(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]", s):
        return s[2:12]
    # Standard 10-character PAN
    if len(s) == 10 and re.match(r"^[A-Z]{5}[0-9]{4}[A-Z]$", s):
        return s
    # 12+ chars starting with 2 digits state code
    if len(s) >= 12 and s[:2].isdigit():
        return s[2:12]
    return s

def find_header_row_and_mapping(ws) -> Tuple[int, Dict[str, int]]:
    """Scans top rows to detect header row and column positions."""
    for r in range(1, min(15, ws.max_row) + 1):
        mapping = {}
        for c in range(1, ws.max_column + 1):
            cell_txt = clean_str(ws.cell(r, c).value).lower()
            if not cell_txt:
                continue
            for key, aliases in HEADER_ALIASES.items():
                if key not in mapping:
                    for alias in aliases:
                        if cell_txt == alias or (len(alias) > 3 and alias in cell_txt):
                            mapping[key] = c
                            break
        if "gstin" in mapping and ("invoice_no" in mapping or "taxable_value" in mapping):
            return r, mapping
    raise ValueError("Could not find table headers with GSTIN/PAN and Invoice No.")

def parse_excel_records(file_bytes: bytes, preferred_sheet_keywords: List[str] = None) -> List[Dict[str, Any]]:
    """
    Parses Excel records, cleans strings/numbers, and extracts PAN from GSTIN/PAN column.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    target_sheet = None
    if preferred_sheet_keywords:
        for kw in preferred_sheet_keywords:
            for sname in wb.sheetnames:
                if kw.lower() in sname.lower():
                    target_sheet = wb[sname]
                    break
            if target_sheet:
                break
    if not target_sheet:
        target_sheet = wb[wb.sheetnames[0]]

    header_row, col_map = find_header_row_and_mapping(target_sheet)

    records = []
    for r in range(header_row + 1, target_sheet.max_row + 1):
        gstin_raw = clean_str(target_sheet.cell(r, col_map.get("gstin", 3)).value).upper()
        inv_no = clean_invoice_no(target_sheet.cell(r, col_map.get("invoice_no", 6)).value)

        # Ignore total rows or blank rows
        if not gstin_raw or len(gstin_raw) < 5 or gstin_raw.startswith("TOTAL") or "GRAND" in gstin_raw:
            continue
        if not inv_no:
            continue

        pan = extract_pan(gstin_raw)
        reco_key = f"{pan}_{inv_no}"

        igst = clean_num(target_sheet.cell(r, col_map.get("igst")).value) if "igst" in col_map else 0.0
        cgst = clean_num(target_sheet.cell(r, col_map.get("cgst")).value) if "cgst" in col_map else 0.0
        sgst = clean_num(target_sheet.cell(r, col_map.get("sgst")).value) if "sgst" in col_map else 0.0
        cess = clean_num(target_sheet.cell(r, col_map.get("cess")).value) if "cess" in col_map else 0.0

        records.append({
            "period": target_sheet.cell(r, col_map.get("period")).value if "period" in col_map else "",
            "trade_name": clean_str(target_sheet.cell(r, col_map.get("trade_name")).value) if "trade_name" in col_map else "",
            "gstin_raw": gstin_raw,
            "pan": pan,
            "gstin": pan,  # use PAN as party identifier so GST and PAN match identically
            "inv_type": clean_str(target_sheet.cell(r, col_map.get("inv_type")).value) if "inv_type" in col_map else "Regular",
            "reverse_charge": clean_str(target_sheet.cell(r, col_map.get("reverse_charge")).value) if "reverse_charge" in col_map else "No",
            "invoice_no": inv_no,
            "reco_key": reco_key,
            "invoice_date": target_sheet.cell(r, col_map.get("invoice_date")).value if "invoice_date" in col_map else None,
            "invoice_value": clean_num(target_sheet.cell(r, col_map.get("invoice_value")).value) if "invoice_value" in col_map else 0.0,
            "taxable_value": clean_num(target_sheet.cell(r, col_map.get("taxable_value")).value) if "taxable_value" in col_map else 0.0,
            "tax_rate": clean_num(target_sheet.cell(r, col_map.get("tax_rate")).value) if "tax_rate" in col_map else 0.0,
            "igst": igst,
            "cgst": cgst,
            "sgst": sgst,
            "cess": cess,
            "total_tax": round(igst + cgst + sgst + cess, 2)
        })

    return records
