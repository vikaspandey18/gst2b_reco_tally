import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from typing import Dict, Any

# Clean styling definitions
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=10, bold=True)
FONT_REGULAR = Font(name="Calibri", size=10)

FILL_DARK_BLUE = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
FILL_MED_BLUE = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FILL_WARNING = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FILL_DIFF = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
FILL_MATCH = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
NUM_FORMAT_CURRENCY = "#,##0.00"

def style_header_row(ws, row_idx: int, col_count: int, fill=FILL_DARK_BLUE):
    for c in range(1, col_count + 1):
        cell = ws.cell(row_idx, c)
        cell.font = FONT_HEADER
        cell.fill = fill
        cell.alignment = ALIGN_CENTER

def generate_reco_excel(reco_data: Dict[str, Any], output_path: Path) -> Path:
    """
    Builds the 5-sheet reconciliation Excel workbook:
    1. GST + Invocie No
    2. GSTR-2B
    3. Book
    4. GST Wise Reco
    5. Duplicate Invoices
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_inv = wb.create_sheet(title="GST + Invocie No")
    ws_2b = wb.create_sheet(title="GSTR-2B")
    ws_book = wb.create_sheet(title="Book")
    ws_gst_reco = wb.create_sheet(title="GST Wise Reco")
    ws_dups = wb.create_sheet(title="Duplicate Invoices")

    gst_summaries = reco_data["gst_wise_summary"]
    inv_items = reco_data["invoice_reco_items"]
    two_b_records = reco_data["two_b_records"]
    book_records = reco_data["book_records"]
    duplicate_records = reco_data.get("duplicate_invoices", [])

    num_inv_rows = len(inv_items)
    num_gst_rows = len(gst_summaries)
    num_2b_rows = len(two_b_records)
    num_book_rows = len(book_records)

    # -------------------------------------------------------------
    # 1. SHEET: GSTR-2B
    # -------------------------------------------------------------
    headers_2b = ["Period", "Trade/Legal name", "GSTIN/PAN", "Invoice Type", "Reverse Charge", "Invoice No", "Key", "Invoice Date", "Invoice Value", "Taxable Value", "Tax Rate", "IGST", "CGST", "SGST", "Cess", "Final Remarks "]
    for c, h in enumerate(headers_2b, 1):
        ws_2b.cell(3, c, h)
    style_header_row(ws_2b, 3, len(headers_2b), FILL_DARK_BLUE)

    for idx, r in enumerate(two_b_records, start=4):
        ws_2b.cell(idx, 1, r.get("period", ""))
        ws_2b.cell(idx, 2, r.get("trade_name", ""))
        ws_2b.cell(idx, 3, r.get("gstin_raw", r.get("pan", "")))
        ws_2b.cell(idx, 4, r.get("inv_type", ""))
        ws_2b.cell(idx, 5, r.get("reverse_charge", ""))
        ws_2b.cell(idx, 6, r.get("invoice_no", ""))
        ws_2b.cell(idx, 7, r.get("reco_key", ""))
        ws_2b.cell(idx, 8, r.get("invoice_date", ""))
        ws_2b.cell(idx, 9, r.get("invoice_value", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_2b.cell(idx, 10, r.get("taxable_value", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_2b.cell(idx, 11, r.get("tax_rate", 0))
        ws_2b.cell(idx, 12, r.get("igst", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_2b.cell(idx, 13, r.get("cgst", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_2b.cell(idx, 14, r.get("sgst", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_2b.cell(idx, 15, r.get("cess", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_2b.cell(idx, 16, f"=VLOOKUP(G{idx},'GST + Invocie No'!$B$5:$Y${max(5, num_inv_rows + 5)},24,0)")

    if num_2b_rows > 0:
        for c in [12, 13, 14]:
            col_letter = get_column_letter(c)
            ws_2b.cell(1, c, f"=SUM({col_letter}4:{col_letter}{num_2b_rows + 3})").font = FONT_BOLD

    # -------------------------------------------------------------
    # 2. SHEET: Book
    # -------------------------------------------------------------
    headers_bk = ["Period", "Trade/Legal name", "GSTIN/PAN", "Invoice Type", "Reverse Charge", "Invoice No", "Key", "Invoice Date", "Invoice Value", "Taxable Value", "Tax Rate", "IGST", "CGST", "SGST", "Cess", "Final Report "]
    for c, h in enumerate(headers_bk, 1):
        ws_book.cell(3, c, h)
    style_header_row(ws_book, 3, len(headers_bk), FILL_MED_BLUE)

    for idx, r in enumerate(book_records, start=4):
        ws_book.cell(idx, 1, r.get("period", ""))
        ws_book.cell(idx, 2, r.get("trade_name", ""))
        ws_book.cell(idx, 3, r.get("gstin_raw", r.get("pan", "")))
        ws_book.cell(idx, 4, r.get("inv_type", ""))
        ws_book.cell(idx, 5, r.get("reverse_charge", ""))
        ws_book.cell(idx, 6, r.get("invoice_no", ""))
        ws_book.cell(idx, 7, r.get("reco_key", ""))
        ws_book.cell(idx, 8, r.get("invoice_date", ""))
        ws_book.cell(idx, 9, r.get("invoice_value", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_book.cell(idx, 10, r.get("taxable_value", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_book.cell(idx, 11, r.get("tax_rate", 0))
        ws_book.cell(idx, 12, r.get("igst", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_book.cell(idx, 13, r.get("cgst", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_book.cell(idx, 14, r.get("sgst", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_book.cell(idx, 15, r.get("cess", 0.0)).number_format = NUM_FORMAT_CURRENCY
        ws_book.cell(idx, 16, f"=VLOOKUP(G{idx},'GST + Invocie No'!$B$5:$Y${max(5, num_inv_rows + 5)},24,0)")

    # -------------------------------------------------------------
    # 3. SHEET: GST Wise Reco
    # -------------------------------------------------------------
    ws_gst_reco.cell(1, 1, "GST / PAN Wise Reconciliation Summary").font = FONT_BOLD
    ws_gst_reco.cell(2, 3, "As per Book ").font = FONT_BOLD
    ws_gst_reco.cell(2, 9, "As per 2B").font = FONT_BOLD
    ws_gst_reco.cell(2, 15, "Diff Book vs 2B").font = FONT_BOLD

    headers_reco = [
        "PAN / GST No", "0", "Book Taxable Amount", "Book IGST", "Book CGST", "Book SGST", "Cess", "0",
        "2B Taxable Amount", "2B IGST", "2B CGST", "2B SGST", "Cess", "0",
        "Taxable Amount", "IGST", "CGST", "SGST", "Cess", "Status"
    ]
    for c, h in enumerate(headers_reco, 1):
        ws_gst_reco.cell(3, c, h)
    style_header_row(ws_gst_reco, 3, len(headers_reco), FILL_DARK_BLUE)

    for idx, g in enumerate(gst_summaries, start=4):
        ws_gst_reco.cell(idx, 1, g["pan"])
        ws_gst_reco.cell(idx, 2, 0)
        # SUMIF matching PAN substring in Book/2B GSTIN/PAN column C
        ws_gst_reco.cell(idx, 3, f'=SUMIF(Book!$C:$C,"*"&$A{idx}&"*",Book!J:J)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 4, f'=SUMIF(Book!$C:$C,"*"&$A{idx}&"*",Book!L:L)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 5, f'=SUMIF(Book!$C:$C,"*"&$A{idx}&"*",Book!M:M)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 6, f'=SUMIF(Book!$C:$C,"*"&$A{idx}&"*",Book!N:N)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 7, f'=SUMIF(Book!$C:$C,"*"&$A{idx}&"*",Book!O:O)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 8, 0)

        ws_gst_reco.cell(idx, 9, f'=SUMIF(\'GSTR-2B\'!$C:$C,"*"&$A{idx}&"*",\'GSTR-2B\'!J:J)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 10, f'=SUMIF(\'GSTR-2B\'!$C:$C,"*"&$A{idx}&"*",\'GSTR-2B\'!L:L)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 11, f'=SUMIF(\'GSTR-2B\'!$C:$C,"*"&$A{idx}&"*",\'GSTR-2B\'!M:M)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 12, f'=SUMIF(\'GSTR-2B\'!$C:$C,"*"&$A{idx}&"*",\'GSTR-2B\'!N:N)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 13, f'=SUMIF(\'GSTR-2B\'!$C:$C,"*"&$A{idx}&"*",\'GSTR-2B\'!O:O)').number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 14, 0)

        ws_gst_reco.cell(idx, 15, f"=+C{idx}-I{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 16, f"=+D{idx}-J{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 17, f"=+E{idx}-K{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 18, f"=+F{idx}-L{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_gst_reco.cell(idx, 19, f"=+G{idx}-M{idx}").number_format = NUM_FORMAT_CURRENCY

        stat_cell = ws_gst_reco.cell(idx, 20, g["status"])
        stat_cell.fill = FILL_MATCH if g["status"] == "Book And 2B Match" else FILL_DIFF

    # -------------------------------------------------------------
    # 4. SHEET: GST + Invocie No
    # -------------------------------------------------------------
    ws_inv.cell(1, 24, "Book And 2B Match-Invocies No Mis-Match").font = FONT_BOLD
    ws_inv.cell(3, 5, "As per Book ").font = FONT_BOLD
    ws_inv.cell(3, 11, "As per 2B").font = FONT_BOLD
    ws_inv.cell(3, 17, "Diff Book vs 2B").font = FONT_BOLD

    headers_inv = [
        "PAN / GST No", "GST + Invoices No", "Counts", "0", "Book Taxable Amount", "Book IGST", "Book CGST", "Book SGST", "Cess", "0",
        "2B Taxable Amount", "2B IGST", "2B CGST", "2B SGST", "Cess", "0",
        "Taxable Amount", "IGST", "CGST", "SGST", "Cess", "0",
        "GST Wise Status", "Status", "Final Remarks "
    ]
    for c, h in enumerate(headers_inv, 1):
        ws_inv.cell(4, c, h)
    style_header_row(ws_inv, 4, len(headers_inv), FILL_DARK_BLUE)

    for idx, item in enumerate(inv_items, start=5):
        ws_inv.cell(idx, 1, item["pan"])
        ws_inv.cell(idx, 2, item["reco_key"])
        ws_inv.cell(idx, 4, 0)
        # SUMIF from Book by reco_key
        ws_inv.cell(idx, 5, f"=SUMIF(Book!$G:$G,$B{idx},Book!J:J)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 6, f"=SUMIF(Book!$G:$G,$B{idx},Book!L:L)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 7, f"=SUMIF(Book!$G:$G,$B{idx},Book!M:M)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 8, f"=SUMIF(Book!$G:$G,$B{idx},Book!N:N)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 9, f"=SUMIF(Book!$G:$G,$B{idx},Book!O:O)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 10, 0)

        # SUMIF from 2B by reco_key
        ws_inv.cell(idx, 11, f"=SUMIF('GSTR-2B'!$G:$G,$B{idx},'GSTR-2B'!J:J)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 12, f"=SUMIF('GSTR-2B'!$G:$G,$B{idx},'GSTR-2B'!L:L)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 13, f"=SUMIF('GSTR-2B'!$G:$G,$B{idx},'GSTR-2B'!M:M)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 14, f"=SUMIF('GSTR-2B'!$G:$G,$B{idx},'GSTR-2B'!N:N)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 15, f"=SUMIF('GSTR-2B'!$G:$G,$B{idx},'GSTR-2B'!O:O)").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 16, 0)

        ws_inv.cell(idx, 17, f"=+E{idx}-K{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 18, f"=+F{idx}-L{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 19, f"=+G{idx}-M{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 20, f"=+H{idx}-N{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 21, f"=+I{idx}-O{idx}").number_format = NUM_FORMAT_CURRENCY
        ws_inv.cell(idx, 22, 0)

        gst_end_row = max(4, num_gst_rows + 3)
        ws_inv.cell(idx, 23, f"=VLOOKUP(A{idx},'GST Wise Reco'!$A$4:$T${gst_end_row},20,0)")
        ws_inv.cell(idx, 24, item["status_x"])
        final_cell = ws_inv.cell(idx, 25, item["final_remarks"])
        final_cell.fill = FILL_MATCH if item["final_remarks"] == "Book And 2B Match" else FILL_DIFF

    # Row 2 Subtotals on Sheet 1
    last_inv_row = max(5, num_inv_rows + 4)
    for c in range(17, 22):
        col_letter = get_column_letter(c)
        ws_inv.cell(2, c, f"=SUBTOTAL(9,{col_letter}5:{col_letter}{last_inv_row})").font = FONT_BOLD
        ws_inv.cell(2, c).number_format = NUM_FORMAT_CURRENCY

    # Bottom Grand Totals
    tot_row = last_inv_row + 2
    ws_inv.cell(tot_row, 2, "Grand Total").font = FONT_BOLD
    for c in [5, 6, 7, 8, 9, 11, 12, 13, 14, 15, 17, 18, 19, 20, 21]:
        col_letter = get_column_letter(c)
        c_cell = ws_inv.cell(tot_row, c, f"=SUM({col_letter}5:{col_letter}{last_inv_row})")
        c_cell.font = FONT_BOLD
        c_cell.number_format = NUM_FORMAT_CURRENCY

    # -------------------------------------------------------------
    # 5. SHEET: Duplicate Invoices (New Sheet)
    # -------------------------------------------------------------
    ws_dups.cell(1, 1, "Duplicate Invoices Report").font = FONT_BOLD
    headers_dup = [
        "Source File", "Period", "Trade/Legal name", "GSTIN / PAN", "Invoice No",
        "Invoice Date", "Invoice Value", "Taxable Value", "Tax Rate", "IGST",
        "CGST", "SGST", "Cess", "Remarks"
    ]
    for c, h in enumerate(headers_dup, 1):
        ws_dups.cell(3, c, h)
    style_header_row(ws_dups, 3, len(headers_dup), FILL_DARK_BLUE)

    if duplicate_records:
        for idx, d in enumerate(duplicate_records, start=4):
            ws_dups.cell(idx, 1, d.get("source", ""))
            ws_dups.cell(idx, 2, d.get("period", ""))
            ws_dups.cell(idx, 3, d.get("trade_name", ""))
            ws_dups.cell(idx, 4, d.get("gstin_raw", d.get("pan", "")))
            ws_dups.cell(idx, 5, d.get("invoice_no", ""))
            ws_dups.cell(idx, 6, d.get("invoice_date", ""))
            ws_dups.cell(idx, 7, d.get("invoice_value", 0.0)).number_format = NUM_FORMAT_CURRENCY
            ws_dups.cell(idx, 8, d.get("taxable_value", 0.0)).number_format = NUM_FORMAT_CURRENCY
            ws_dups.cell(idx, 9, d.get("tax_rate", 0))
            ws_dups.cell(idx, 10, d.get("igst", 0.0)).number_format = NUM_FORMAT_CURRENCY
            ws_dups.cell(idx, 11, d.get("cgst", 0.0)).number_format = NUM_FORMAT_CURRENCY
            ws_dups.cell(idx, 12, d.get("sgst", 0.0)).number_format = NUM_FORMAT_CURRENCY
            ws_dups.cell(idx, 13, d.get("cess", 0.0)).number_format = NUM_FORMAT_CURRENCY
            ws_dups.cell(idx, 14, d.get("remark", "Duplicate invoice"))
            ws_dups.cell(idx, 14).fill = FILL_WARNING
    else:
        ws_dups.cell(4, 1, "No duplicate invoices found in GSTR-2B or Book.")

    # Column widths and grid lines for all sheets
    for sheet in [ws_inv, ws_2b, ws_book, ws_gst_reco, ws_dups]:
        sheet.views.sheetView[0].showGridLines = True
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:15]:
                if cell.value and not str(cell.value).startswith("="):
                    max_len = max(max_len, len(str(cell.value)))
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path
